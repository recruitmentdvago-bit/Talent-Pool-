import io
import time
from pathlib import Path

import pytest

pytest.importorskip("flask")
from sourcer import runner  # noqa: E402
from sourcer.ask import parse_ask  # noqa: E402
from sourcer.web.app import JOB, create_app  # noqa: E402

H = {"X-Sourcer-UI": "1"}

JD_FORM = {"title": "Data Analyst", "seniority": "mid", "min_years": 2, "education_min": "bachelors",
           "openings": 2, "must_have": "SQL, Power BI, Excel", "nice_have": "Python",
           "hard_requirements": "SQL", "industries": "banking"}


def signed_in(app):
    """A test client already logged in as the built-in admin."""
    c = app.test_client()
    assert c.post("/api/login", json={"username": "admin", "password": "admin123"}, headers=H).status_code == 200
    return c


@pytest.fixture
def client(sandbox):
    JOB.reset()
    app = create_app()
    app.config["TESTING"] = True
    return signed_in(app)


def wait_job(client, seconds=60):
    end = time.time() + seconds
    while time.time() < end:
        j = client.get("/api/job").get_json()
        if j["state"] not in {"running", "stopping"}:
            return j
        time.sleep(0.2)
    raise AssertionError("job did not finish")


# ---- the natural-language request ---------------------------------------- #
@pytest.mark.parametrize("text,role,score", [
    ("data analyst manager ke liye resumes nikaldo jiski score 70 ya 75% se above ho", "data analyst manager", 70),
    ("Give me data analysts and data engineers above 80%", "data analyst/data engineer", 80),
    ("DA-001 ke 65 se upar wale", "DA-001", 65),
    ("show all candidates", None, None),
    ("senior data engineer chahiye 90 percent", "senior data engineer", 90),
])
def test_parse_ask(text, role, score):
    r = parse_ask(text)
    assert r["role"] == role and r["min_score"] == score


# ---- local-only protection ------------------------------------------------ #
def test_only_localhost_may_talk_to_it(client):
    assert client.get("/api/overview").status_code == 200
    assert client.get("/api/overview", headers={"Host": "attacker.example"}).status_code == 403
    # 127.0.0.1 is allowed; the session cookie belongs to "localhost", so this host is simply not signed in
    assert client.get("/api/overview", headers={"Host": "127.0.0.1:8765"}).status_code == 401


def test_writes_need_the_app_header(client):
    assert client.post("/api/job/stop", json={}).status_code == 403
    assert client.post("/api/job/stop", json={}, headers=H).status_code == 200
    assert client.delete("/api/jds/X").status_code == 403


# ---- vacancies ------------------------------------------------------------ #
def test_vacancy_create_edit_archive(client, sandbox):
    r = client.post("/api/jds", json=JD_FORM, headers=H).get_json()
    assert r["ok"] and r["code"] == "V-001"
    files = list((sandbox / "jds").glob("V-001-*.md"))
    assert len(files) == 1
    jds = client.get("/api/jds").get_json()
    assert [j["code"] for j in jds] == ["V-001"]
    assert jds[0]["must_have"] == ["SQL", "Power BI", "Excel"] and jds[0]["version"] == 1

    edit = dict(JD_FORM, code="V-001", min_years=5)
    client.post("/api/jds", json=edit, headers=H)
    jds = client.get("/api/jds").get_json()
    assert jds[0]["min_years"] == 5 and jds[0]["version"] == 2
    assert len(list((sandbox / "jds").glob("V-001*.md"))) == 1, "an edit must replace the file"

    assert client.delete("/api/jds/V-001", headers=H).status_code == 200
    assert client.get("/api/jds").get_json() == []
    assert list((sandbox / "jds" / "_archived").glob("V-001*"))
    client.post("/api/run", json={"kind": "rescore"}, headers=H)
    wait_job(client)
    assert client.get("/api/jds").get_json() == [], "an archived vacancy must stay archived"


def test_second_vacancy_gets_the_next_code(client):
    client.post("/api/jds", json=JD_FORM, headers=H)
    r = client.post("/api/jds", json=dict(JD_FORM, title="Data Engineer"), headers=H).get_json()
    assert r["code"] == "V-002"


def test_vacancy_needs_a_title(client):
    assert client.post("/api/jds", json={"title": " "}, headers=H).status_code == 400


def test_pasted_jd_text_is_read_into_fields(client):
    r = client.post("/api/jds/preview", headers=H, json={"text":
        "Business Intelligence Analyst\n\nRequirements\n- 3-5 years of experience\n- SQL and Power BI\n"
        "- Bachelors degree\n\nNice to have\n- Python"}).get_json()
    assert r["min_years"] == 3 and r["max_years"] == 5
    assert "sql" in [s.lower() for s in r["must_have"]] and "python" in [s.lower() for s in r["nice_have"]]
    assert client.post("/api/jds/preview", json={"text": ""}, headers=H).status_code == 400


# ---- intake, candidates, ask, export ------------------------------------- #
def _pdf_bytes(sandbox, strong_analyst):
    from sourcer.pdfstore import render_profile_pdf
    p = sandbox / "ayesha.pdf"
    render_profile_pdf(strong_analyst, p)
    return p.read_bytes()


def test_upload_process_ask_show_and_export(client, sandbox, strong_analyst):
    client.post("/api/jds", json=JD_FORM, headers=H)
    up = client.post("/api/upload", headers=H, content_type="multipart/form-data", data={
        "files": [(io.BytesIO(_pdf_bytes(sandbox, strong_analyst)), "ayesha.pdf"),
                  (io.BytesIO(b"x"), "notes.txt")]}).get_json()
    assert up["pdfs"] == 1 and up["skipped"] == ["notes.txt"]
    assert client.get("/api/overview").get_json()["inbox"]["pdfs"] == 1

    assert client.post("/api/run", json={"kind": "files"}, headers=H).status_code == 200
    job = wait_job(client)
    assert job["state"] == "done", job["log"]
    assert client.get("/api/overview").get_json()["candidates"] == 1

    a = client.post("/api/ask", headers=H, json={"text": "data analyst ke liye 75% se upar"}).get_json()
    assert a["matched"][0]["code"] == "V-001" and a["min_score"] == 75
    assert [r["name"] for r in a["rows"]] == ["Ayesha Khan"]
    cid = a["rows"][0]["id"]
    assert a["rows"][0]["has_pdf"] is True

    none = client.post("/api/ask", headers=H, json={"text": "astronaut 70"}).get_json()
    assert none["rows"] == [] and "No open vacancy" in none["message"]

    d = client.get(f"/api/candidates/{cid}").get_json()
    assert d["scores"][0]["breakdown"]["must_have_skills"]["raw"] == 1.0
    assert d["experiences"] and d["has_pdf"]

    pdf = client.get(f"/api/candidates/{cid}/pdf")
    assert pdf.status_code == 200 and pdf.data[:4] == b"%PDF"

    assert client.post(f"/api/candidates/{cid}/status", json={"status": "shortlisted"}, headers=H).status_code == 200
    assert client.post(f"/api/candidates/{cid}/status", json={"status": "bogus"}, headers=H).status_code == 400
    shortlisted = client.get("/api/candidates?status=shortlisted&min_score=0").get_json()
    assert [r["id"] for r in shortlisted] == [cid]

    for fmt, magic in (("xlsx", b"PK"), ("csv", b"candidate_id"), ("zip", b"PK")):
        r = client.get(f"/api/export?format={fmt}&min_score=0")
        assert r.status_code == 200 and r.data[:len(magic)] == magic, fmt
    assert client.get("/api/export?format=xlsx&min_score=0&status=rejected").status_code == 404


def test_pdf_endpoint_will_not_serve_files_outside_the_library(client, sandbox, strong_analyst):
    from sourcer import db, store
    secret = sandbox / "secret.pdf"
    secret.write_bytes(b"%PDF-1.4 secret")
    strong_analyst.profile_pdf_path = str(secret)
    with db.session() as c:
        cid, _ = store.upsert(c, strong_analyst)
    assert client.get(f"/api/candidates/{cid}/pdf").status_code == 404


# ---- settings -------------------------------------------------------------- #
def test_settings_are_saved_to_the_ui_file_and_clamped(client, sandbox):
    r = client.post("/api/settings", headers=H, json={
        "linkedin": {"daily_profile_cap": 9999, "window_start": "9:30", "window_end": "nonsense",
                     "session_minutes_max": 1, "location_filter": "bogus",
                     "geo_urn": "https://www.linkedin.com/search/results/people/?geoUrn=%5B%22105343366%22%5D&origin=x"},
        "threshold": 70})
    assert r.status_code == 200
    s = client.get("/api/settings").get_json()
    assert s["linkedin"]["daily_profile_cap"] == 150
    assert s["linkedin"]["window_start"] == "9:30" and s["linkedin"]["window_end"] == "19:00"
    assert s["linkedin"]["session_minutes_max"] == 5
    assert s["linkedin"]["location_filter"] == "metro"
    assert s["linkedin"]["geo_urn"] == "%5B%22105343366%22%5D"
    assert s["threshold"] == 70
    assert (sandbox / "settings.ui.yaml").exists()
    assert "geo_urn" not in (sandbox / "settings.yaml").read_text()


def test_only_one_job_runs_at_a_time(client):
    import threading
    gate = threading.Event()
    assert JOB.start("x", lambda j: gate.wait(5)) is True
    assert client.post("/api/run", json={"kind": "files"}, headers=H).status_code == 409
    gate.set()
    wait_job(client)


def test_linkedin_start_is_refused_until_logged_in(client):
    r = client.post("/api/run", json={"kind": "linkedin"}, headers=H)
    assert r.status_code == 400 and "logged in" in r.get_json()["error"]


def test_cooldown_is_visible_and_clearable(client):
    from sourcer import db
    from sourcer.pacing import start_cooldown
    with db.session() as c:
        start_cooldown(c, "test warning", hours=5)
    o = client.get("/api/overview").get_json()["linkedin"]
    assert o["cooldown_until"] and o["cooldown_reason"] == "test warning"
    client.post("/api/linkedin/clear-cooldown", json={}, headers=H)
    assert client.get("/api/overview").get_json()["linkedin"]["cooldown_until"] is None


def test_only_web_addresses_become_links(client):
    from sourcer import db, store
    from sourcer.models import Candidate, Experience
    with db.session() as c:
        cid, _ = store.upsert(c, Candidate(
            full_name="Evil Link", linkedin_url="javascript:alert(1)", location_raw="Karachi, Pakistan",
            experiences=[Experience(title="Data Analyst", company="X", start_date="Jan 2020", is_current=True)]))
        cid2, _ = store.upsert(c, Candidate(full_name="Fine Link", linkedin_url="https://www.linkedin.com/in/fine-1/"))
    assert client.get(f"/api/candidates/{cid}").get_json()["linkedin_url"] == ""
    assert client.get(f"/api/candidates/{cid2}").get_json()["linkedin_url"].startswith("https://")


def _docx(paragraphs, bullets=()):
    """A minimal real .docx in memory."""
    import zipfile
    body = "".join(f"<w:p><w:r><w:t>{t}</w:t></w:r></w:p>" for t in paragraphs)
    body += "".join(f"<w:p><w:pPr><w:numPr/></w:pPr><w:r><w:t>{t}</w:t></w:r></w:p>" for t in bullets)
    xml = ('<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="http://schemas.openxmlformats.org/'
           f'wordprocessingml/2006/main"><w:body>{body}</w:body></w:document>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/document.xml", xml)
    return buf.getvalue()


def test_vacancy_from_a_word_file_fills_the_form(client):
    data = _docx(["Business Intelligence Analyst", "Requirements"],
                 bullets=["3-5 years of experience", "Strong SQL and Power BI", "Bachelors degree"])
    r = client.post("/api/jds/from-files", headers=H, content_type="multipart/form-data",
                    data={"mode": "extract", "files": (io.BytesIO(data), "jd.docx")}).get_json()
    assert "Business Intelligence Analyst" in r["text"] and "- Strong SQL" in r["text"]
    p = client.post("/api/jds/preview", json={"text": r["text"]}, headers=H).get_json()
    assert p["min_years"] == 3 and "sql" in [x.lower() for x in p["must_have"]]


def test_many_jd_files_become_many_vacancies(client, sandbox):
    from sourcer.pdfstore import render_profile_pdf  # any real PDF writer works for a JD too
    docs = [("a.docx", _docx(["Data Analyst", "Requirements"], bullets=["2+ years", "SQL and Excel"])),
            ("b.txt", b"Data Engineer\n\nRequirements\n- 5+ years\n- Python, SQL, Airflow\n"),
            ("c.doc", b"\xd0\xcf\x11\xe0 old binary word"),
            ("d.xyz", b"?")]
    r = client.post("/api/jds/from-files", headers=H, content_type="multipart/form-data",
                    data={"mode": "create", "files": [(io.BytesIO(b), n) for n, b in docs]}).get_json()
    assert [c["title"] for c in r["created"]] == ["Data Analyst", "Data Engineer"]
    assert {f["file"] for f in r["failed"]} == {"c.doc", "d.xyz"}
    assert any(".docx" in f["error"] for f in r["failed"] if f["file"] == "c.doc")
    jds = client.get("/api/jds").get_json()
    assert {j["code"] for j in jds} == {"V-001", "V-002"} and all(j["must_have"] for j in jds)


def test_vacancies_survive_a_restart(client, sandbox):
    from sourcer import config
    client.post("/api/jds", json=JD_FORM, headers=H)
    client.post("/api/jds", json=dict(JD_FORM, title="Data Engineer"), headers=H)
    JOB.reset()
    config.cfg(reload=True)                       # a fresh start of the app
    again = signed_in(create_app())
    assert {j["title"] for j in again.get("/api/jds").get_json()} == {"Data Analyst", "Data Engineer"}
    assert len(list((sandbox / "jds").glob("V-*.md"))) == 2


def test_a_typed_percentage_beats_the_slider(client, sandbox, strong_analyst):
    """'80%' in the sentence is used, whatever the slider says; with no number the slider value is used."""
    client.post("/api/jds", json=JD_FORM, headers=H)
    a = client.post("/api/ask", headers=H, json={"text": "data analyst 80% se upar"}).get_json()
    assert a["min_score"] == 80
    b = client.post("/api/ask", headers=H, json={"text": "data analyst ke resumes dikhao"}).get_json()
    assert b["min_score"] == 75, "no number typed -> the configured default"


def test_with_no_number_typed_the_slider_value_is_used(client):
    client.post("/api/jds", json=JD_FORM, headers=H)
    r = client.post("/api/ask", headers=H, json={"text": "data analyst chahiye", "default_min": "70"}).get_json()
    assert r["min_score"] == 70
    r = client.post("/api/ask", headers=H, json={"text": "data analyst 85% se upar", "default_min": "70"}).get_json()
    assert r["min_score"] == 85, "a typed number must beat the slider"


# ---- org structure: departments, HODs, and who sees what ------------------- #
def _hod(app, name):
    c = app.test_client()
    assert c.post("/api/login", json={"username": name, "password": "secret1"}, headers=H).status_code == 200
    return c


def test_hods_only_see_their_departments_vacancies_and_candidates(sandbox, strong_analyst):
    JOB.reset()
    app = create_app()
    admin = signed_in(app)
    admin.post("/api/jds", json=JD_FORM, headers=H)                                    # Data Analyst
    admin.post("/api/jds", json={**JD_FORM, "title": "Accountant", "must_have": "Excel"}, headers=H)
    admin.post("/api/departments/import", headers=H, json={"text": "Technology: Data Analyst\nFinance: Accountant"})
    for u in ("tech.hod", "fin.hod", "lost.hod"):
        assert admin.post("/api/users", headers=H, json={"username": u, "password": "secret1", "role": "hod"}).status_code == 200
    org = admin.get("/api/org").get_json()
    ids = {u["username"]: u["id"] for u in org["hods"]}
    dept = {d["name"]: d for d in org["departments"]}
    assert [v["title"] for v in dept["Technology"]["vacancies"]] == ["Data Analyst"] and org["unassigned"] == []
    for name, user in (("Technology", "tech.hod"), ("Finance", "fin.hod")):
        r = admin.put(f"/api/departments/{dept[name]['id']}", headers=H,
                      json={"name": name, "roles": dept[name]["roles"], "hods": [ids[user]]})
        assert r.status_code == 200

    pdf = _pdf_bytes(sandbox, strong_analyst)
    admin.post("/api/upload", headers=H, content_type="multipart/form-data", data={"files": [(io.BytesIO(pdf), "a.pdf")]})
    admin.post("/api/run", json={"kind": "files"}, headers=H)
    assert wait_job(admin)["state"] == "done"
    cid = admin.get("/api/candidates?min_score=0").get_json()[0]["id"]

    tech, fin, lost = _hod(app, "tech.hod"), _hod(app, "fin.hod"), _hod(app, "lost.hod")
    assert [j["title"] for j in tech.get("/api/jds").get_json()] == ["Data Analyst"]
    assert [j["title"] for j in fin.get("/api/jds").get_json()] == ["Accountant"]
    assert lost.get("/api/jds").get_json() == []
    assert {r["vacancy_title"] for r in tech.get("/api/candidates?min_score=0").get_json()} == {"Data Analyst"}
    assert {r["vacancy_title"] for r in fin.get("/api/candidates?min_score=0").get_json()} == {"Accountant"}
    assert lost.get("/api/candidates?min_score=0").get_json() == []
    assert [s["title"] for s in fin.get(f"/api/candidates/{cid}").get_json()["scores"]] == ["Accountant"]
    assert lost.get(f"/api/candidates/{cid}").status_code == 404
    assert lost.get(f"/api/candidates/{cid}/pdf").status_code == 404
    assert lost.post(f"/api/candidates/{cid}/status", json={"status": "rejected"}, headers=H).status_code == 404
    assert lost.get("/api/export?format=csv&min_score=0").status_code == 404
    assert tech.post(f"/api/candidates/{cid}/status", json={"status": "reviewed"}, headers=H).status_code == 200

    for method, url in (("post", "/api/jds"), ("post", "/api/run"), ("post", "/api/upload"), ("get", "/api/org"),
                        ("get", "/api/users"), ("get", "/api/settings"), ("post", "/api/departments"),
                        ("put", "/api/jds/V-001/department")):
        assert getattr(tech, method)(url, json={}, headers=H).status_code == 403, url
    ov = tech.get("/api/overview").get_json()
    assert ov["linkedin"] is None and ov["job"]["log"] == []


def test_assigning_a_department_beats_role_matching_and_survives_edits(client):
    client.post("/api/jds", json=JD_FORM, headers=H)
    client.post("/api/departments", json={"name": "Technology", "roles": "Data Analyst"}, headers=H)
    client.post("/api/departments", json={"name": "Finance", "roles": ""}, headers=H)
    assert client.get("/api/jds").get_json()[0]["department"] == "Technology"
    assert client.put("/api/jds/V-001/department", json={"department": "Nowhere"}, headers=H).status_code == 400
    assert client.put("/api/jds/V-001/department", json={"department": "finance"}, headers=H).status_code == 200
    j = client.get("/api/jds").get_json()[0]
    assert (j["department"], j["department_how"]) == ("Finance", "assigned")
    client.post("/api/jds", json={**JD_FORM, "code": "V-001", "openings": 3}, headers=H)       # edit without a department
    assert client.get("/api/jds").get_json()[0]["department"] == "Finance"
    client.post("/api/jds", json={**JD_FORM, "code": "V-001", "department": ""}, headers=H)    # form clears it
    assert client.get("/api/jds").get_json()[0]["department_how"] == "role"


def test_renaming_a_department_keeps_its_hods_and_vacancies(client):
    client.post("/api/jds", json=JD_FORM, headers=H)
    client.post("/api/users", json={"username": "ali", "password": "secret1", "role": "hod"}, headers=H)
    client.post("/api/departments", json={"name": "Tech", "roles": "Data Analyst"}, headers=H)
    org = client.get("/api/org").get_json()
    did, uid = org["departments"][0]["id"], org["hods"][0]["id"]
    client.put(f"/api/departments/{did}", json={"name": "Tech", "roles": "Data Analyst", "hods": [uid]}, headers=H)
    client.put("/api/jds/V-001/department", json={"department": "Tech"}, headers=H)
    client.put(f"/api/departments/{did}", json={"name": "Technology", "roles": "Data Analyst", "hods": [uid]}, headers=H)
    assert client.get("/api/jds").get_json()[0]["department"] == "Technology"
    ali = [u for u in client.get("/api/users").get_json() if u["username"] == "ali"][0]
    assert ali["departments"] == ["Technology"] and ali["visible_vacancies"] == 1
    assert client.post("/api/departments", json={"name": "technology"}, headers=H).status_code == 409
    client.delete(f"/api/departments/{did}", headers=H)
    ali = [u for u in client.get("/api/users").get_json() if u["username"] == "ali"][0]
    assert ali["departments"] == [] and ali["visible_vacancies"] == 0


def test_default_admin_password_is_flagged_until_changed(client):
    me = client.get("/api/me").get_json()
    assert me["default_password"] is True and me["version"]
    r = client.put(f"/api/users/{me['id']}", json={"password": "a-better-one"}, headers=H)
    assert r.status_code == 200
    assert client.get("/api/me").get_json()["default_password"] is False
    assert client.put(f"/api/users/{me['id']}", json={"role": "hod"}, headers=H).status_code == 400   # last admin


# ---- central rules: preferred companies, search titles, templates, bulk ---- #
def test_company_lists_reach_matching_vacancies_and_change_scores(client, sandbox, strong_analyst):
    client.post("/api/jds", json=JD_FORM, headers=H)                                     # Data Analyst V-001
    client.post("/api/jds", json={**JD_FORM, "title": "Accountant", "must_have": "Excel"}, headers=H)
    client.post("/api/departments", json={"name": "Technology", "roles": "Data Analyst"}, headers=H)
    assert client.post("/api/employer-groups", json={"name": "x"}, headers=H).status_code == 400   # needs companies
    r = client.post("/api/employer-groups", json={"name": "Tech employers", "companies": "Systems Limited, Engro",
                                                   "departments": ["Technology"]}, headers=H)
    assert r.status_code == 200
    client.post("/api/employer-groups", json={"name": "Everyone", "companies": ["Jazz"]}, headers=H)   # no scope = all
    jds = {j["title"]: j for j in client.get("/api/jds").get_json()}
    assert sorted(jds["Data Analyst"]["effective_employers"]) == ["Engro", "Jazz", "Systems Limited"]
    assert jds["Accountant"]["effective_employers"] == ["Jazz"]
    assert jds["Data Analyst"]["preferred_employers"] == []                               # own list untouched
    g = client.get("/api/employer-groups").get_json()["groups"]
    assert {x["name"]: len(x["vacancies"]) for x in g} == {"Tech employers": 1, "Everyone": 2}
    # the loader the scorer and LinkedIn use sees the same thing
    from sourcer import db, jd as jdmod
    with db.session() as c:
        by = {j.title: j for _, j in jdmod.load_all(c)}
        assert sorted(by["Data Analyst"].preferred_employers) == ["Engro", "Jazz", "Systems Limited"]
        assert jdmod.load_all(c, effective=False)[0][1].preferred_employers == []


def test_search_titles_are_generated_and_managed_centrally(client):
    client.post("/api/jds", json={**JD_FORM, "title": "Senior Data Analyst"}, headers=H)
    rules_ = client.get("/api/title-rules").get_json()["rules"]
    assert [r["role"] for r in rules_] == ["Senior Data Analyst"] and rules_[0]["auto"] is True
    assert "Business Intelligence Analyst" in rules_[0]["titles"]
    assert not any("head" in t.lower() or "manager" in t.lower() for t in rules_[0]["titles"])
    # a shorter role rule also reaches the longer vacancy title
    client.post("/api/title-rules", json={"role": "Data Analyst", "titles": "Decision Analyst"}, headers=H)
    j = client.get("/api/jds").get_json()[0]
    assert "Decision Analyst" in j["effective_titles"] and "Decision Analyst" in j["titles_from_rules"]
    from sourcer import db, jd as jdmod
    from sourcer.sources.linkedin_browser import search_keywords
    with db.session() as c:
        jd = jdmod.load_all(c)[0][1]
    q = search_keywords(jd)
    assert q.index('"Senior Data Analyst"') < q.index('"Decision Analyst"')
    assert client.post("/api/title-rules", json={"role": "data analyst"}, headers=H).status_code == 409
    eff = client.post("/api/jds/effective", json={"title": "Data Analyst"}, headers=H).get_json()
    assert "Decision Analyst" in eff["titles"] and eff["suggested_titles"]
    assert client.post("/api/title-rules/generate", json={}, headers=H).get_json()["created"] == 0


def test_word_template_round_trips_into_every_field(client):
    r = client.get("/api/templates/vacancy.docx")
    assert r.status_code == 200 and r.data[:2] == b"PK"
    from sourcer.jd import docx_to_text
    text = docx_to_text(r.data)
    assert "Job title:" in text and "Must-have skills:" in text and "Job description:" in text
    filled = text.replace("Job title:", "Job title: Pharmacy Manager").replace("Department:", "Department: Pharmacy Operations") \
        .replace("Level:", "Level: senior").replace("Min years:", "Min years: 5") \
        .replace("Must-have skills:", "Must-have skills: Inventory, Dispensing") \
        .replace("Job description:", "Job description:\nLead the store team and the dispensing counter.")
    pv = client.post("/api/jds/preview", json={"text": filled}, headers=H).get_json()
    assert (pv["title"], pv["department"], pv["seniority"], pv["min_years"]) == ("Pharmacy Manager", "Pharmacy Operations", "senior", 5.0)
    assert pv["must_have"] == ["Inventory", "Dispensing"] and pv["from_template"] is True
    assert pv["body"].startswith("Lead the store team")


def test_excel_template_and_bulk_upload(client):
    client.post("/api/departments", json={"name": "Technology", "roles": "Data Analyst"}, headers=H)
    r = client.get("/api/templates/vacancies.xlsx")
    assert r.status_code == 200
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Vacancies", "Example", "How to fill"]
    ws = wb["Vacancies"]
    heads = [c.value for c in ws[1]]
    assert heads[0] == "Job title *" and "Preferred employers" in heads and "Extra titles" not in heads
    col = {h: i + 1 for i, h in enumerate(heads)}
    rows = [{"Job title *": "Data Analyst", "Department": "Technology", "Level": "mid", "Min years": 3,
             "Must-have skills": "SQL, Power BI", "Openings": 2},
            {"Job title *": "Accountant", "Must-have skills": "Excel", "Minimum education": "masters"},
            {"Department": "Technology"},                                   # no title -> reported, not fatal
            {"Job title *": "Data Analyst"}]                              # duplicate of row 1 -> skipped
    for n, row in enumerate(rows, start=2):
        for h, v in row.items():
            ws.cell(row=n, column=col[h], value=v)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    res = client.post("/api/jds/from-sheet", headers=H, content_type="multipart/form-data",
                      data={"file": (buf, "mine.xlsx")}).get_json()
    assert [c["title"] for c in res["created"]] == ["Data Analyst", "Accountant"]
    assert [f["row"] for f in res["failed"]] == [4] and [s["row"] for s in res["skipped"]] == [5]
    jds = {j["title"]: j for j in client.get("/api/jds").get_json()}
    assert jds["Data Analyst"]["department"] == "Technology" and jds["Data Analyst"]["openings"] == 2
    assert jds["Accountant"]["education_min"] == "masters"
    assert {r["role"] for r in client.get("/api/title-rules").get_json()["rules"]} == {"Data Analyst", "Accountant"}
    csv_bytes = io.BytesIO(b"Job title,Min years\nStore Manager,4\n")
    res = client.post("/api/jds/from-sheet", headers=H, content_type="multipart/form-data",
                      data={"file": (csv_bytes, "x.csv")}).get_json()
    assert res["created"][0]["title"] == "Store Manager"
    bad = client.post("/api/jds/from-sheet", headers=H, content_type="multipart/form-data",
                      data={"file": (io.BytesIO(b"a,b\n1,2\n"), "x.csv")})
    assert bad.status_code == 400


def test_hods_cannot_reach_the_rule_screens(sandbox):
    JOB.reset()
    app = create_app()
    admin = signed_in(app)
    admin.post("/api/users", headers=H, json={"username": "some.hod", "password": "secret1", "role": "hod"})
    hod = _hod(app, "some.hod")
    for method, url in (("get", "/api/employer-groups"), ("get", "/api/title-rules"), ("get", "/api/templates/vacancies.xlsx"),
                        ("get", "/api/templates/vacancy.docx"), ("post", "/api/jds/from-sheet"), ("post", "/api/jds/effective"),
                        ("post", "/api/employer-groups"), ("post", "/api/title-rules")):
        assert getattr(hod, method)(url, headers=H).status_code == 403, url


def test_titles_come_from_the_jd_not_from_typing(client, monkeypatch):
    from sourcer import rules
    # no key -> the taxonomy plus the roles whose core skills the JD asks for
    client.post("/api/jds", json={**JD_FORM, "must_have": "SQL, Power BI, Tableau"}, headers=H)
    r = client.get("/api/title-rules").get_json()["rules"][0]
    assert r["source"] == "jd" and "BI Analyst" in r["titles"] and "BI Developer" in r["titles"]
    # a key -> Claude's reading of the JD comes first, and is not redone while the JD is unchanged
    calls = []
    monkeypatch.setattr(rules, "_ai_titles", lambda t, text, sk: calls.append(t) or ["Decision Scientist", "Insights Lead"])
    client.post("/api/title-rules/generate", json={"force": True}, headers=H)
    r = client.get("/api/title-rules").get_json()["rules"][0]
    assert r["source"] == "ai" and r["titles"][:2] == ["Decision Scientist", "Insights Lead"] and len(calls) == 1
    client.post("/api/title-rules/generate", json={}, headers=H)
    assert len(calls) == 1
    # editing the JD re-reads it; a rule you edited yourself is left alone
    client.post("/api/jds", json={**JD_FORM, "code": "V-001", "must_have": "SQL, Power BI, Tableau", "text": "Now with a new description."}, headers=H)
    assert len(calls) == 2
    rid = client.get("/api/title-rules").get_json()["rules"][0]["id"]
    client.put(f"/api/title-rules/{rid}", json={"role": "Data Analyst", "titles": "My Own Title"}, headers=H)
    client.post("/api/jds", json={**JD_FORM, "code": "V-001", "text": "Changed again."}, headers=H)
    r = client.get("/api/title-rules").get_json()["rules"][0]
    assert r["titles"] == ["My Own Title"] and r["auto"] is False and len(calls) == 2
    # the company mode is a validated setting
    assert client.post("/api/settings", json={"linkedin": {"company_mode": "bogus"}}, headers=H).status_code == 200
    assert client.get("/api/settings").get_json()["linkedin"]["company_mode"] == "priority_first"
    client.post("/api/settings", json={"linkedin": {"company_mode": "only"}}, headers=H)
    assert client.get("/api/settings").get_json()["linkedin"]["company_mode"] == "only"


def test_bulk_download_has_every_matching_resume_even_without_a_stored_pdf(client, sandbox, strong_analyst):
    import zipfile
    from sourcer import db, store
    client.post("/api/jds", json=JD_FORM, headers=H)
    strong_analyst.profile_pdf_path = ""                       # e.g. a LinkedIn profile or spreadsheet row: no PDF yet
    with db.session() as c:
        store.upsert(c, strong_analyst)
        from sourcer import runner
        runner.score_all(c)
    r = client.get("/api/export?format=zip&role=V-001&min_score=0")
    assert r.status_code == 200 and r.data[:2] == b"PK"
    assert "V-001" in r.headers["Content-Disposition"] and "resumes" in r.headers["Content-Disposition"]
    z = zipfile.ZipFile(io.BytesIO(r.data))
    names = z.namelist()
    assert any(n.endswith("Ayesha_Khan.pdf") for n in names), names
    assert "candidates.xlsx" in names and "00_cover_sheet.pdf" in names and "MISSING_PDFS.txt" not in names
    pdf = z.read(next(n for n in names if n.endswith("Ayesha_Khan.pdf")))
    assert pdf[:4] == b"%PDF"
